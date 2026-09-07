"""Word(.docx) 산출 — 합본 진단 결과 보고서.

PDF 는 그대로 제출하는 확정본이고, Word 는 **담당자가 고쳐 쓰는 원본**이다. 진단 결과에
발주처 양식·기관 표지·현장에서 확인한 정오탐 의견을 얹어 최종 산출물로 만드는 일이
실제 진단의 대부분이라, 같은 내용을 편집 가능한 형태로도 낸다.

구성·문안은 pdf.combined_report 와 같다(표지·개정이력·목차·1~5장·부록). 색은
report/style.py 의 흑백 팔레트를 공유한다.

목차는 Word 필드로 넣는다 — 여는 시점에 Word 가 쪽번호를 계산해야 편집 후에도 맞는다.
"""
from __future__ import annotations

import datetime as _dt
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from ..i18n import (DEFAULT_REM_EN, REMEDIATION_EN, SEV_EN, STEP_LABEL,
                    STEP_LABEL_EN, tr)
from . import style as S
from .pdf import (SEV_KR, SEV_ORDER, REMEDIATION, _CRITERIA, _CRITERIA_EN,
                  _DEFAULT_REM, _PRIORITY, VULN_EXAMPLE, _rule_key)

_FONT_KO = "맑은 고딕"
_FONT_MONO = "D2Coding"       # 없으면 Word 가 대체 폰트를 쓴다
_DETAIL_CAP = 200             # 상세 기술 상한 (PDF 와 동일)


def _hex(c: str) -> RGBColor:
    return RGBColor.from_string(c.lstrip("#").upper())


def _shade(cell, hexcolor: str) -> None:
    """셀 배경. python-docx 에 API 가 없어 w:shd 를 직접 넣는다."""
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor.lstrip("#").upper())
    cell._tc.get_or_add_tcPr().append(el)


def _set_font(run, *, size=10, bold=False, color=S.INK, mono=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _hex(color)
    name = _FONT_MONO if mono else _FONT_KO
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    return run


def _para(doc, text="", *, size=10, bold=False, color=S.INK, align=None,
          space_after=4, mono=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    if text:
        _set_font(p.add_run(text), size=size, bold=bold, color=color, mono=mono)
    return p


def _heading(doc, text, level=1):
    """제목. Word 기본 Heading 스타일을 쓰되 글자색만 검정으로 맞춘다(목차 필드가 이걸 읽는다)."""
    h = doc.add_heading(level=level)
    h.paragraph_format.space_before = Pt(14 if level == 1 else 10)
    h.paragraph_format.space_after = Pt(6 if level == 1 else 3)
    _set_font(h.add_run(text), size=15 if level == 1 else 11.5, bold=True, color=S.INK)
    return h


def _toc_field(doc, label: str) -> None:
    """목차 자리에 TOC 필드를 심는다. Word 가 열 때(또는 F9) 쪽번호를 계산한다."""
    p = doc.add_paragraph()
    r = p.add_run()
    fld = OxmlElement("w:fldChar"); fld.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-2" \h \z \u'
    sep = OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"), "separate")
    hint = OxmlElement("w:t"); hint.text = label
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    for el in (fld, instr, sep, hint, end):
        r._r.append(el)


def _table(doc, rows, widths_cm, *, header=True, sizes=9, aligns=None):
    """표 하나. rows[0] 을 머리글로 본다."""
    t = doc.add_table(rows=0, cols=len(widths_cm))
    t.style = "Table Grid"
    t.autofit = False
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            cell = cells[ci]
            cell.width = Cm(widths_cm[ci])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.space_before = Pt(1)
            if aligns and aligns[ci] is not None:
                p.alignment = aligns[ci]
            txt, color, bold = val if isinstance(val, tuple) else (val, S.INK, False)
            _set_font(p.add_run(str(txt)), size=sizes,
                      bold=bold or (header and ri == 0), color=color)
            if header and ri == 0:
                _shade(cell, S.FILL_HEAD)
    for row in t.rows:                       # 열 너비는 셀마다 다시 지정해야 먹는다
        for ci, w in enumerate(widths_cm):
            row.cells[ci].width = Cm(w)
    return t


def _kv(doc, rows, w0=4.5, w1=12.5):
    data = [(k, v) for k, v in rows]
    t = _table(doc, data, [w0, w1], header=False, sizes=9.5)
    for r in t.rows:
        _shade(r.cells[0], S.FILL_HEAD)
        for run in r.cells[0].paragraphs[0].runs:
            run.font.bold = True
    return t


CENTER = WD_ALIGN_PARAGRAPH.CENTER


def combined_report(scan, path, author: str = "CPGuard", lang: str = "ko",
                    meta: dict | None = None,
                    standards: list[str] | str | None = None) -> None:
    """합본 진단 결과 보고서(Word). 인자는 pdf.combined_report 와 같다."""
    from .. import standards as _sm

    en = lang == "en"
    T = lambda s: tr(s, lang)                       # noqa: E731
    SEV = SEV_EN if en else SEV_KR
    REM = REMEDIATION_EN if en else REMEDIATION
    DFT = DEFAULT_REM_EN if en else _DEFAULT_REM
    CRIT = _CRITERIA_EN if en else _CRITERIA
    meta = meta or {}
    author = meta.get("author") or author
    version = meta.get("version") or "1.0"
    findings = scan.findings
    counts = scan.severity_counts
    total = len(findings)
    project = scan.project or Path(scan.name).stem
    today = _dt.date.today().strftime("%Y-%m-%d")
    order = {s: i for i, s in enumerate(SEV_ORDER)}

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = _FONT_KO
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_KO)

    # ── 표지 ──
    _para(doc, "", space_after=60)
    _para(doc, "SOURCE CODE SECURITY ASSESSMENT", size=11, color=S.MUTED, align=CENTER)
    _para(doc, project, size=22, bold=True, align=CENTER, space_after=2)
    _para(doc, T("소스코드 취약점 진단 결과 보고서"), size=18, bold=True, align=CENTER)
    _para(doc, T("정적 보안 진단 · CPGuard"), size=12, color=S.INK_SOFT, align=CENTER, space_after=24)

    cover = [(T("프로젝트"), project), (T("대상"), scan.name)]
    if meta.get("client"):
        cover.append((T("발주처/고객"), meta["client"]))
    if meta.get("org"):
        cover.append((T("수행 기관/회사"), meta["org"]))
    cover += [(T("분석 도구"), T("CPGuard (CPG 기반 taint 분석)")),
              (T("분석 일시"), scan.created_at.strftime("%Y-%m-%d %H:%M"))]
    if meta.get("period"):
        cover.append((T("진단 수행 기간"), meta["period"]))
    cover += [(T("소스 파일 수"), str(scan.file_count)), (T("총 이슈"), str(total)),
              (T("작성일"), today), (T("작성자"), author)]
    if meta.get("tester"):
        cover.append((T("진단 담당자"), meta["tester"]))
    cover.append((T("보고서 버전"), version))
    _kv(doc, cover)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ── 문서 개정 이력 ── (목차에는 넣지 않는다 — Heading 스타일을 쓰지 않음)
    _para(doc, T("문서 개정 이력"), size=12, bold=True, space_after=6)
    _table(doc, [[T("버전"), T("일자"), T("내용"), T("작성")],
                 [version, today, T("최초 작성"), author]],
           [2.0, 3.0, 9.6, 2.8], aligns=[CENTER, CENTER, None, CENTER])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ── 목차 ──
    _heading(doc, T("목차"), 1)
    _toc_field(doc, T("목차를 갱신하려면 이 영역을 선택하고 F9 를 누르세요."))
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ── 1. 진단 개요 ──
    _heading(doc, T("1. 진단 개요"), 1)
    _heading(doc, T("1.1 진단 배경 및 목적"), 2)
    _para(doc, T("본 보고서는 대상 소스코드에 대해 CPGuard 정적 분석(데이터 흐름 taint 분석 + 패턴 점검)을 "
                 "수행하여 도출한 보안약점과 그 조치 방안을 기술한다. 각 취약점은 CWE·OWASP 기준으로 분류하고, "
                 "위험도에 따라 조치 우선순위를 제시한다."), space_after=10)

    langs = ", ".join(sorted({Path(f["file"]).suffix.lstrip(".").lower()
                              for f in findings if f.get("file")}) or ["-"])
    _heading(doc, T("1.2 진단 대상 범위"), 2)
    _kv(doc, [(T("프로젝트"), project), (T("대상"), scan.name),
              (T("대상 파일 수"), str(scan.file_count)), (T("탐지 이슈 수"), str(total)),
              (T("포함 확장자"), langs)])

    _heading(doc, T("1.3 진단 방법 및 기준"), 2)
    method = [(T("진단 방식"), T("정적 분석 — 데이터 흐름(taint) + 패턴 점검")),
              (T("분석 도구"), "CPGuard"),
              (T("진단 일시"), scan.created_at.strftime("%Y-%m-%d %H:%M"))]
    _kv(doc, method)
    if scan.integrity_note:
        _para(doc, "* " + scan.integrity_note, size=8.5, color=S.INK_SOFT)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ── 2. 진단 결과 요약 ──
    _heading(doc, T("2. 진단 결과 요약"), 1)
    rows = [[T("위험도"), T("건수"), T("비율")]]
    for s in SEV_ORDER:
        n = counts.get(s, 0)
        rows.append([SEV.get(s, s), str(n), f"{(n / total * 100 if total else 0):.1f}%"])
    rows.append([(T("합계"), S.INK, True), (str(total), S.INK, True), ("100.0%", S.INK, True)])
    _table(doc, rows, [5.0, 4.0, 4.0], aligns=[None, CENTER, CENTER])

    _para(doc, "", space_after=8)
    _heading(doc, T("취약점 유형(CWE) 상위"), 2)
    cwe_c = Counter(f.get("cwe") or "-" for f in findings)
    seen: dict = {}
    for f in findings:
        seen.setdefault(f.get("cwe") or "-", f.get("rule_id"))
    rows = [["CWE", T("규칙 예"), T("개수")]]
    for cwe, n in cwe_c.most_common(12):
        rows.append([cwe, seen.get(cwe, ""), str(n)])
    _table(doc, rows, [3.6, 11.8, 2.0], aligns=[None, None, CENTER])

    # 진단 이력
    Scan = scan.__class__
    runs = (list(Scan.objects.filter(project=scan.project).order_by("-created_at"))
            if scan.project else [scan])
    _para(doc, "", space_after=8)
    _heading(doc, T("진단 이력"), 2)
    _para(doc, (f"This project has been assessed {len(runs)} time(s). Runs are listed latest "
                f"first; 'New'/'Resolved' are relative to the immediately preceding run."
                if en else
                f"이 프로젝트는 총 {len(runs)}회 진단되었다. 최근 회차부터 나열하며, "
                f"신규/해결은 직전 회차 대비 증감이다(해결된 취약점은 해당 회차 건수에서 이미 제외됨)."),
          size=9.5, space_after=6)
    rows = [[T("회차"), T("일시"), T("탐지"), T("신규"), T("해결")] + [SEV.get(s, s) for s in SEV_ORDER]]
    cur_idx = next((i for i, r in enumerate(runs) if r.pk == scan.pk), 0)
    for i, r in enumerate(runs):
        mark = " ◀" if r.pk == scan.pk else ""
        rows.append([f"{len(runs) - i}{mark}", r.created_at.strftime("%Y-%m-%d %H:%M"),
                     str(r.finding_count),
                     f"+{r.new_count}" if r.new_count else "0",
                     f"-{r.resolved_count}" if r.resolved_count else "0",
                     str(r.sev_critical), str(r.sev_high), str(r.sev_medium),
                     str(r.sev_low), str(r.sev_info)])
    ht = _table(doc, rows, [1.5, 3.0, 1.6, 1.6, 1.6] + [1.62] * 5,
                sizes=8.5, aligns=[CENTER] * 10)
    for cell in ht.rows[cur_idx + 1].cells:      # 이번 회차 강조
        _shade(cell, S.FILL_ZEBRA)
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ── 3. 진단 항목 ──
    _heading(doc, T("3. 진단 항목"), 1)
    ids = [standards] if isinstance(standards, str) else list(standards or [])
    stds = [s for s in (_sm.get(i) for i in ids) if s]
    cwe_counts = Counter(c for f in findings
                         if (c := (f.get("cwe") or "").strip().upper()))

    if not stds:
        _para(doc, T("이번 진단에서 탐지된 점검 항목(규칙)과 분류·건수는 다음과 같다."), space_after=6)
        by_rule: dict = {}
        for f in findings:
            d = by_rule.setdefault(f["rule_id"], {"cwe": f.get("cwe", ""), "n": 0})
            d["n"] += 1
        rows = [[T("점검 항목"), "CWE", T("탐지")]]
        for r, d in sorted(by_rule.items(), key=lambda x: -x[1]["n"]):
            rows.append([r, d["cwe"] or "-", str(d["n"])])
        _table(doc, rows, [12.4, 3.2, 1.8], aligns=[None, None, CENTER])
    else:
        avail = _sm.rule_cwes()
        V = _sm.VERDICT_EN if en else _sm.VERDICT_KO
        for std in stds:
            crows = _sm.coverage(std, cwe_counts, avail)
            nv = sum(1 for r in crows if r["verdict"] == _sm.VIOLATED)
            nn = sum(1 for r in crows if r["verdict"] == _sm.NOT_COVERED)
            _heading(doc, std.name_en if en else std.name, 2)
            intro = (f"Assessed against {std.source_for(lang)}. {nv} of the {len(crows)} weaknesses below "
                     f"were found."
                     if en else
                     f"점검 기준은 {std.source_for(lang)}이며, 아래 {len(crows)}개 보안약점 중 "
                     f"{nv}개 항목에서 위반이 확인되었다.")
            if nn:
                intro += (f" {nn} items are outside static analysis and are marked "
                          f"'{V[_sm.NOT_COVERED]}' rather than passing."
                          if en else
                          f" 이 중 {nn}개 항목은 정적 분석으로 확인할 수 있는 규칙이 없어 "
                          f"양호가 아니라 ‘{V[_sm.NOT_COVERED]}’ 으로 표기한다.")
            _para(doc, intro, size=9.5, space_after=6)

            code = std.show_code
            head = ([T("유형")] + ([T("코드")] if code else [])
                    + [T("보안약점"), "CWE", T("판정"), T("탐지")])
            rows = [head]
            for r in crows:
                cw = (", ".join(r["cwes"][:3]) + ("…" if len(r["cwes"]) > 3 else "")) or "-"
                dim = S.MUTED if r["verdict"] == _sm.NOT_COVERED else S.INK
                vb = r["verdict"] == _sm.VIOLATED
                rows.append([(r["group_en"] if en else r["group"], dim, False)]
                            + ([(r["code"], dim, False)] if code else [])
                            + [(r["name_en"] if en else r["name"], dim, False),
                               (cw, dim, False), (V[r["verdict"]], dim, vb),
                               (str(r["n"]) if r["n"] else "-", dim, False)])
            widths = ([2.9] + ([1.5] if code else [])
                      + [(5.6 if code else 7.0), 3.0, 2.4, 1.4])
            _table(doc, rows, widths, sizes=8,
                   aligns=[None] * (len(head) - 3) + [CENTER, CENTER, CENTER])
            _para(doc, "", space_after=8)

        um = {}
        for std in stds:
            um.update(_sm.unmapped(std, cwe_counts))
        um = {c: n for c, n in um.items() if not any(s.item_for(c) for s in stds)}
        if um:
            lead = ("Findings outside the selected standards: "
                    if en else "선택한 기준의 점검항목에 매핑되지 않은 탐지: ")
            _para(doc, lead + ", ".join(f"{c}({n})" for c, n in
                                        sorted(um.items(), key=lambda kv: -kv[1])),
                  size=8.5, color=S.INK_SOFT)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ── 4. 상세 진단 결과 ──
    _heading(doc, T("4. 상세 진단 결과"), 1)
    ordered = sorted(findings, key=lambda x: (order.get(x["severity"], 9),
                                              x.get("file", ""), x.get("line", 0)))
    shown = ordered[:_DETAIL_CAP]
    if total > _DETAIL_CAP:
        _para(doc, (f"* Showing the top {_DETAIL_CAP} findings by severity; see the analysis "
                    f"sheet (xlsx) for all {total}."
                    if en else
                    f"* 위험도 상위 {_DETAIL_CAP}건을 상세 기술하며, 전체 {total}건은 "
                    f"분석목록표(xlsx)를 참조한다."),
              size=8.5, color=S.INK_SOFT)
    slabel = STEP_LABEL_EN if en else STEP_LABEL
    for i, f in enumerate(shown, 1):
        _finding_block(doc, i, f, SEV, REM, DFT, T, en, slabel)

    # ── 5. 종합 의견 ──
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    _heading(doc, T("5. 종합 의견"), 1)
    ch = counts.get("critical", 0) + counts.get("high", 0)
    _para(doc, (f"A total of {total} security weaknesses were identified, of which {ch} are "
                f"Critical/High severity requiring immediate action. Prioritize Critical/High "
                f"items, then apply input validation, output encoding, secret separation and "
                f"safe algorithms per the remediation for each type."
                if en else
                f"총 {total}건의 보안약점이 도출되었으며, 이 중 즉시 조치가 필요한 매우위험·위험 "
                f"등급이 {ch}건이다. 매우위험·위험 항목을 우선 조치하고, 유형별 조치 방안에 따라 "
                f"입력 검증·출력 인코딩·비밀정보 분리·안전한 알고리즘 적용을 권고한다."))

    # ── 부록 A ──
    _para(doc, "", space_after=10)
    _heading(doc, T("부록 A. 위험도 판정 기준"), 1)
    rows = [[T("판정"), T("기준"), T("조치 우선순위")]]
    for s in SEV_ORDER:
        rows.append([(SEV.get(s, s), S.SEV_INK[s], True), CRIT[s], T(_PRIORITY[s])])
    _table(doc, rows, [2.8, 11.8, 2.8], sizes=9)

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))


def _finding_block(doc, idx, f, SEV, REM, DFT, T, en, slabel) -> None:
    """취약점 1건 — 제목 줄 + 항목별 상세 표. PDF 의 카드와 같은 정보를 담는다."""
    sev = f["severity"]
    cwe = f.get("cwe") or ""
    owasp = f.get("owasp") or ""
    tail = f"  ({cwe}{' · ' + owasp if owasp else ''})" if cwe or owasp else ""
    p = _para(doc, "", space_after=2)
    _set_font(p.add_run(f"[{SEV.get(sev, sev)}] {idx}. {f['rule_id']}{tail}"),
              size=10.5, bold=True, color=S.SEV_INK.get(sev, S.INK))

    rk = _rule_key(f["rule_id"])
    rem = REM.get(rk, DFT)
    rows = [(T("대상"), f"{f['file']}:{f['line']}"),
            (T("설명"), T(f.get("message", "")))]
    steps = f.get("steps") or []
    if steps:
        flow = "\n".join(
            f"[{slabel.get(s.get('kind', ''), s.get('kind', ''))}] "
            f"{s.get('file', '')}:{s.get('line', '')}  {(s.get('code') or '').strip()[:160]}"
            for s in steps[:12])
        rows.append((T("데이터 흐름"), flow))
    rows += [(T("영향"), rem[1]), (T("조치 방안"), rem[2])]
    if bad := VULN_EXAMPLE.get(rk):
        rows.append((T("취약한 코드 예시"), bad))
    if rem[3]:
        rows.append((T("안전한 코드 예시"), rem[3]))
    ref = cwe + (f" · OWASP {owasp}" if owasp else "")
    if ref:
        rows.append((T("참고"), ref))

    t = _table(doc, [[k, v] for k, v in rows], [3.0, 14.4], header=False, sizes=9)
    for r in t.rows:
        _shade(r.cells[0], S.FILL_HEAD)
        for run in r.cells[0].paragraphs[0].runs:
            run.font.bold = True
    # 데이터 흐름·코드 예시는 고정폭으로 (줄 맞춤이 의미를 갖는다)
    for r, (k, _v) in zip(t.rows, rows, strict=True):
        if k in (T("데이터 흐름"), T("취약한 코드 예시"), T("안전한 코드 예시")):
            for run in r.cells[1].paragraphs[0].runs:
                run.font.name = _FONT_MONO
                run.font.size = Pt(8)
                run._element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_MONO)
    _para(doc, "", space_after=8)


# ==================== 합본 진단 결과 보고서 (다중 프로젝트) ====================

def consolidated_report(scans, path, lang: str = "ko", meta: dict | None = None,
                        standards: list[str] | str | None = None) -> None:
    """합본 진단 결과 보고서(Word). 구성·문안은 pdf.consolidated_report 와 같다.

    Word 본이 담당자가 고쳐 쓰는 원본이다 — 발주처 양식·기관 표지·현장 의견을 얹어
    최종 산출물로 만드는 일이 실제 진단의 대부분이다.
    """
    from . import consolidated as C
    from .pdf import _tool_version

    en = lang == "en"
    T = lambda s: tr(s, lang)                       # noqa: E731
    SEV = SEV_EN if en else SEV_KR
    REM = REMEDIATION_EN if en else REMEDIATION
    DFT = DEFAULT_REM_EN if en else _DEFAULT_REM
    meta = meta or {}
    scans = list(scans)
    D = C.build(scans, standards, lang)
    stds = D["standards"]

    version = meta.get("version") or "1.0"
    author = meta.get("author") or "CPGuard"
    today = _dt.date.today().strftime("%Y-%m-%d")
    client = meta.get("client") or ""
    system = meta.get("system") or ""
    title = " ".join(x for x in (client, system) if x) or T("소스코드 취약점 진단")

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = _FONT_KO
    normal.font.size = Pt(10)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), _FONT_KO)

    # ── 표지 ──
    _para(doc, "", space_after=60)
    _para(doc, "SOURCE CODE SECURITY ASSESSMENT", size=11, color=S.MUTED, align=CENTER)
    _para(doc, title, size=22, bold=True, align=CENTER, space_after=2)
    _para(doc, T("소스코드 취약점 진단 결과 보고서"), size=18, bold=True, align=CENTER)
    _para(doc, T("정적 보안 진단 · CPGuard"), size=12, color=S.INK_SOFT, align=CENTER, space_after=24)
    _kv(doc, [(T("발주처/고객"), client or "-"), (T("대상 시스템"), system or "-"),
              (T("대상 프로젝트"), str(D["projects"]) + ("" if en else "개")),
              (T("수행 기관/회사"), meta.get("org") or "-"),
              (T("진단 수행 기간"), meta.get("period") or "-"),
              (T("보고서 버전"), version), (T("작성일"), today), (T("작성자"), author)])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ── 제·개정 이력 ── (Heading 스타일이 아니라 목차에 안 들어간다)
    _para(doc, T("제·개정 이력"), size=12, bold=True, space_after=6)
    _table(doc, [[T("버전"), T("변경일"), T("변경 사유"), T("변경 내용"), T("작성자"), T("비고")],
                 [version, today, T("최초 작성"), T("최초 작성"), author, "-"]],
           [1.6, 2.4, 3.4, 5.6, 2.6, 1.8], sizes=9,
           aligns=[CENTER, CENTER, None, None, CENTER, CENTER])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ── 목차 ── (Heading 을 쓰지 않아 목차가 자기 자신을 담지 않는다)
    _para(doc, T("목 차"), size=15, bold=True, space_after=8)
    _toc_field(doc, T("목차를 갱신하려면 이 영역을 선택하고 F9 를 누르세요."))
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ── 1. 취약점 진단 개요 ──
    _heading(doc, T("1. 취약점 진단 개요"), 1)
    _heading(doc, T("1.1 진단 목적"), 2)
    subject = (client + " " + system).strip() or T("대상 시스템")
    _para(doc, (f"The purpose of this assessment is to identify and remove security weaknesses "
                f"in the source code of {subject} in advance, so that the threats arising from "
                f"those weaknesses are mitigated and the service and its information are "
                f"protected from malicious internal and external attack."
                if en else
                f"{subject}의 소스코드 보안약점을 도출하여 이를 사전에 제거함으로써, 소스코드 "
                f"보안약점으로 인해 발생할 수 있는 위협에 대한 대응방안을 마련하고, 내·외부의 "
                f"악의적인 공격으로부터 서비스 및 정보를 보호하는 것을 목적으로 한다."))
    if stds:
        basis = ", ".join(s.source_for(lang) for s in stds)
        _para(doc, (f"The assessment is performed in accordance with {basis}." if en
                    else f"{basis}의 기준을 준수하여 진단을 수행한다."), space_after=10)

    _heading(doc, T("1.2 점검 수행 일정"), 2)
    _kv(doc, [(T("진단 수행 기간"), meta.get("period") or "-"),
              (T("보고서 작성일"), today),
              (T("대상 프로젝트 수"), str(D["projects"]) + ("" if en else "개"))])

    _heading(doc, T("1.3 점검 도구"), 2)
    tool = meta.get("tool") or ("CPGuard " + _tool_version()).strip()
    _table(doc, [[T("진단 도구명"), T("용도"), T("비고")],
                 [tool, T("소스코드의 데이터 흐름(taint)과 위험 패턴을 정적으로 분석하여 "
                          "보안약점을 검출하는 정적 분석 도구"), T("CPG 기반")]],
           [4.0, 10.6, 2.8], sizes=9)

    _heading(doc, T("1.4 점검 수행 인원"), 2)
    _table(doc, [[T("이름"), T("직급"), T("이메일"), T("연락처")],
                 [meta.get("tester") or author, meta.get("tester_rank") or "-",
                  meta.get("tester_email") or "-", meta.get("tester_phone") or "-"]],
           [3.4, 2.6, 6.2, 5.2], sizes=9, aligns=[CENTER, CENTER, None, CENTER])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ── 2. 진단 항목 ──
    _heading(doc, T("2. 진단 항목"), 1)
    if not stds:
        _para(doc, T("점검 기준을 지정하지 않아 이번 진단에서 탐지된 규칙 유형을 그대로 싣는다."))
    for std in stds:
        _heading(doc, std.name_en if en else std.name, 2)
        _para(doc, (f"Assessed against {std.source_for(lang)}. The check items are as follows."
                    if en else
                    f"{std.source_for(lang)}에 근거한 진단 항목을 적용한다. 점검 항목은 다음과 같다."),
              size=9.5, space_after=6)
        gs = [g for g in D["groups"] if g["standard"] == (std.name_en if en else std.name)]
        if len(gs) > 1:
            rows = [[T("순번"), T("항목"), T("설명"), T("항목 수")]]
            for i, g in enumerate(gs, 1):
                rows.append([str(i), g["group"], g["desc"], str(g["n"])])
            rows.append([("", S.INK, True), (T("합계"), S.INK, True), "",
                         (str(sum(g["n"] for g in gs)), S.INK, True)])
            _table(doc, rows, [1.3, 3.8, 10.5, 1.8], sizes=8.5,
                   aligns=[CENTER, None, None, CENTER])
        _para(doc, "", space_after=8)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ── 3. 진단 결과 ──
    _heading(doc, T("3. 진단 결과"), 1)
    it, ft = D["initial_total"], D["final_total"]

    def scale_table(rows, total):
        # 빌드 라인은 값이 있을 때만 — 이 컬럼이 생기기 전 스캔은 0 이라 빈 열이 된다.
        show = any(r.lines for r in rows)
        head = ([T("프로젝트"), T("파일 수")] + ([T("빌드 라인")] if show else [])
                + [T("검출")] + [SEV[s] for s in SEV_ORDER])
        data = [head]
        for r in rows:
            data.append([r.name, f"{r.files:,}"] + ([f"{r.lines:,}"] if show else [])
                        + [str(r.total)] + [str(r.sev.get(s, 0)) for s in SEV_ORDER])
        data.append([(T("총 계"), S.INK, True), (f"{total['files']:,}", S.INK, True)]
                    + ([(f"{total['lines']:,}", S.INK, True)] if show else [])
                    + [(str(total["total"]), S.INK, True)]
                    + [(str(total["sev"][s]), S.INK, True) for s in SEV_ORDER])
        widths = ([4.0, 1.7] + ([2.0] if show else []) + [1.5]
                  + [(1.64 if show else 2.04)] * 5)
        _table(doc, data, widths, sizes=8, aligns=[None] + [CENTER] * (len(head) - 1))

    _heading(doc, T("3.1 최초 보안약점 진단 결과"), 2)
    _para(doc, (f"The initial assessment across {D['projects']} project(s) detected "
                f"{it['total']} security weaknesses."
                if en else
                f"총 {D['projects']}개 프로젝트를 대상으로 수행한 최초 소스코드 보안약점 진단 "
                f"결과, {it['total']}건의 보안약점이 검출되었다."), size=9.5, space_after=6)
    scale_table(D["initial"], it)

    _heading(doc, T("3.2 진단 결과 점검"), 2)
    dropped = it["total"] - ft["total"]
    if D["review"]:
        _para(doc, (f"Reviewing the source context, {dropped} of the {it['total']} detections "
                    f"were classified as excluded / false positive / fixed, leaving "
                    f"{ft['total']} items to remediate. The rationale recorded by the assessor "
                    f"is as follows."
                    if en else
                    f"소스 컨텍스트 재확인을 통해 검출 {it['total']}건 중 {dropped}건을 "
                    f"「제외 / 오탐 / 조치완료」로 분류하고, 조치대상 {ft['total']}건을 "
                    f"확정하였다. 항목별 사유와 진단원 의견은 다음과 같다."), size=9.5)
        _para(doc, T("※ 소스 컨텍스트(Source Context) : 검출 지점 전후의 코드 흐름, 호출 관계, "
                     "프레임워크·설정 정보 등 취약점의 실제 성립 여부를 판단하기 위해 참조하는 "
                     "주변 소스코드 정보를 의미한다."), size=8.5, color=S.INK_SOFT, space_after=6)
        rows = [[T("보안약점명"), T("위험도"), T("건수"), T("사유"), T("진단원 의견")]]
        for r in D["review"][:60]:
            rows.append([r["name"], SEV.get(r["severity"], r["severity"]), str(r["n"]),
                         r["reason"], r["opinion"] or "-"])
        rows.append([(T("총 계"), S.INK, True), "", (str(D["review_total"]), S.INK, True), "", ""])
        _table(doc, rows, [3.6, 1.8, 1.4, 1.8, 8.8], sizes=8,
               aligns=[None, CENTER, CENTER, CENTER, None])
    else:
        _para(doc, ("No detection has been reviewed as excluded or a false positive yet, so the "
                    "initial result stands as the final result. Verdicts and opinions recorded "
                    "on the review screen appear in this section."
                    if en else
                    "검토 화면에서 오탐·제외로 판정한 항목이 아직 없어, 최초 진단 결과가 그대로 "
                    "최종 결과가 된다. 검토 화면에서 판정과 의견을 기록하면 이 절에 반영된다."),
              size=9.5)

    _heading(doc, T("3.3 최종 점검 결과"), 2)
    _para(doc, (f"After the review, the initial {it['total']} detections were confirmed as "
                f"{ft['total']} items to remediate."
                if en else
                f"오탐·제외 항목을 정리한 결과, 최초 {it['total']}건에서 최종 {ft['total']}건으로 "
                f"확정되었다."), size=9.5, space_after=6)
    scale_table(D["final"], ft)
    _para(doc, ("Per-project detail follows. Each weakness carries its severity, and where it "
                "maps to a check item of the applied standards that item name is used."
                if en else
                "아래는 프로젝트별 상세 진단 결과이며, 각 항목에 위험도를 부여하고 적용 기준의 "
                "점검항목에 대응되는 경우 그 보안약점명으로 표기하였다."), size=9.5, space_after=8)

    for i, r in enumerate(D["final"], 1):
        _heading(doc, f"3.3.{i} {r.name}", 2)
        _kv(doc, [(T("발주처/고객"), client or "-"), (T("서비스명"), r.name),
                  (T("개발언어"), r.languages), (T("파일 수"), f"{r.files:,}"),
                  (T("빌드 라인 수"), f"{r.lines:,}")], w0=3.4, w1=13.6)
        _para(doc, "", space_after=4)
        rows = [[T("순번"), T("유형"), T("보안약점명"), T("위험도"), T("건수"), T("비고")]]
        if r.weaknesses:
            for j, w in enumerate(r.weaknesses, 1):
                rows.append([str(j), w["group"], w["name"],
                             SEV.get(w["severity"], w["severity"]), str(w["n"]), "-"])
        else:
            rows.append(["1", "-", T("점검 기한 내 발견된 취약점 없음"), "-", "0", "-"])
        rows.append([("", S.INK, True), (T("총 계"), S.INK, True), "", "",
                     (str(r.total), S.INK, True), ""])
        _table(doc, rows, [1.2, 4.0, 6.6, 2.0, 1.6, 2.0], sizes=8,
               aligns=[CENTER, None, None, CENTER, CENTER, CENTER])
        _para(doc, "", space_after=8)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ── 4. 유형별 조치 권고 ──
    _heading(doc, T("4. 유형별 조치 권고"), 1)
    _para(doc, ("Remediation for each weakness type found across the assessed projects. Code "
                "examples for each type are in the separate remediation guide."
                if en else
                "이번 진단에서 도출된 보안약점 유형별 조치 방안은 다음과 같다. 유형별 상세 코드 "
                "예시는 별도의 조치 가이드를 참조한다."), size=9.5, space_after=6)
    seen = Counter(_rule_key(f["rule_id"])
                   for r in D["final"] for f in r.scan.findings)
    rows = [[T("보안약점 유형"), T("영향"), T("조치 방안"), T("건수")]]
    for k, n in seen.most_common(20):
        rem = REM.get(k, DFT)
        rows.append([rem[0], rem[1], rem[2], str(n)])
    _table(doc, rows, [3.4, 6.0, 6.6, 1.4], sizes=8, aligns=[None, None, None, CENTER])

    # ── 5. 종합 의견 ──
    _para(doc, "", space_after=10)
    _heading(doc, T("5. 종합 의견"), 1)
    ch = ft["sev"]["critical"] + ft["sev"]["high"]
    _para(doc, (f"Across {D['projects']} project(s), {it['total']} weaknesses were detected and "
                f"{ft['total']} were confirmed for remediation, of which {ch} are Critical/High "
                f"and require immediate action. Address Critical/High items first, then apply "
                f"input validation, output encoding, secret separation and safe algorithms per "
                f"the remediation for each type."
                if en else
                f"총 {D['projects']}개 프로젝트에서 {it['total']}건이 검출되어 {ft['total']}건이 "
                f"조치대상으로 확정되었으며, 이 중 즉시 조치가 필요한 매우위험·위험 등급이 "
                f"{ch}건이다. 매우위험·위험 항목을 우선 조치하고, 유형별 조치 방안에 따라 "
                f"입력 검증·출력 인코딩·비밀정보 분리·안전한 알고리즘 적용을 권고한다."))

    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
