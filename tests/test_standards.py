"""점검 기준(레퍼런스) 선택 — 매핑·필터·산출물."""
import io
import os
import tempfile
import zipfile

import pytest

os.environ.setdefault("CPGUARD_HOME", tempfile.mkdtemp(prefix="cpguard_std_"))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "cpguard.web.settings")

import django  # noqa: E402

django.setup()
from django.conf import settings  # noqa: E402
from django.core.files.uploadedfile import SimpleUploadedFile  # noqa: E402
from django.core.management import call_command  # noqa: E402
from django.test import Client  # noqa: E402

from cpguard import standards  # noqa: E402

settings.ALLOWED_HOSTS = ["testserver", "127.0.0.1", "localhost"]


@pytest.fixture(scope="module", autouse=True)
def _db():
    call_command("migrate", verbosity=0, interactive=False)


def test_every_rule_cwe_maps_to_a_check_item():
    """규칙이 내보내는 CWE 가 기준 표에서 빠지면 그 탐지는 조용히 '미분류'가 된다.

    기준을 늘리거나 규칙을 추가할 때 여기서 걸리게 둔다. 매핑되지 않는 CWE 가 있는 것
    자체는 정상(모든 기준이 모든 약점을 다루지는 않는다)이지만, 어느 것이 왜 빠졌는지는
    unmapped() 로 산출물에 드러나야 한다."""
    from cpguard.patterns import load_pattern_rules
    from cpguard.taint.spec import load_rules
    cwes = {r.cwe for r in load_rules()} | {r.cwe for r in load_pattern_rules()}
    cwes = {c for c in cwes if c}
    assert cwes, "규칙에 CWE 가 하나도 없다"

    mois = standards.get("mois")
    missing = sorted(c for c in cwes if not mois.item_for(c))
    assert not missing, f"행안부 기준에 매핑 안 된 CWE: {missing}"

    counts = {c: 1 for c in cwes}
    assert not standards.unmapped(mois, counts)


@pytest.mark.parametrize("sid", sorted(standards.STANDARDS))
def test_item_codes_are_unique_and_lookup_works(sid):
    std = standards.get(sid)
    codes = [i.code for i in std.items]
    assert len(codes) == len(set(codes))
    for it in std.items:
        for c in it.cwes:
            assert std.item_for(c) is not None
    assert standards.get("없는기준") is None


def test_coverage_marks_untouched_items_as_pass():
    """걸리지 않은 항목도 표에 남아야 산출물이 '무엇을 점검했는가'의 증빙이 된다."""
    std = standards.get("mois")
    rows = standards.coverage(std, {"CWE-89": 3})
    assert len(rows) == len(std.items)
    hit = [r for r in rows if r["n"]]
    assert len(hit) == 1 and hit[0]["code"] == "sql-injection" and hit[0]["n"] == 3


def _seed(c: Client) -> int:
    from cpguard.web.models import Scan
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("a/i.js", "app.get('/x',(req,res)=>{eval(req.query.q)})")
    c.post("/scan/", {"archive": SimpleUploadedFile("std.zip", buf.getvalue())},
           SERVER_NAME="127.0.0.1")
    return Scan.objects.order_by("-id").first().pk


def test_standard_api_and_item_filter():
    c = Client()
    pk = _seed(c)

    # 기준 없이 부르면 고를 수 있는 목록을 준다
    listing = c.get(f"/scan/{pk}/api/standard", SERVER_NAME="127.0.0.1").json()
    assert {s["id"] for s in listing["standards"]} == set(standards.STANDARDS)

    d = c.get(f"/scan/{pk}/api/standard?std=mois", SERVER_NAME="127.0.0.1").json()
    assert d["total_items"] == len(standards.get("mois").items)
    hit = [i for i in d["items"] if i["n"]]
    assert hit, "eval 탐지가 어느 항목에도 안 걸렸다"

    # 항목 필터가 그 항목 건수와 일치해야 한다
    code = hit[0]["code"]
    got = c.get(f"/scan/{pk}/api/findings?std=mois&item={code}", SERVER_NAME="127.0.0.1").json()
    assert got["total"] == hit[0]["n"]

    # 없는 항목 코드는 0건(전체가 새는 것보다 낫다)
    none = c.get(f"/scan/{pk}/api/findings?std=mois&item=없는항목", SERVER_NAME="127.0.0.1").json()
    assert none["total"] == 0


def test_deliverables_carry_the_chosen_standard():
    c = Client()
    pk = _seed(c)

    import openpyxl
    xl = c.get(f"/scan/{pk}/export.xlsx?std=mois", SERVER_NAME="127.0.0.1")
    assert xl.status_code == 200
    wb = openpyxl.load_workbook(io.BytesIO(xl.content))
    assert "점검항목 결과" in wb.sheetnames
    # 14컬럼 분석목록표 형식은 건드리지 않는다(제출 형식 고정)
    assert wb["분석목록표"].max_column == 14

    plain = openpyxl.load_workbook(io.BytesIO(
        c.get(f"/scan/{pk}/export.xlsx", SERVER_NAME="127.0.0.1").content))
    assert "점검항목 결과" not in plain.sheetnames

    from pypdf import PdfReader
    pdf = c.get(f"/scan/{pk}/report.pdf?std=mois", SERVER_NAME="127.0.0.1")
    assert pdf.status_code == 200
    txt = "".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(pdf.content)).pages)
    assert "행정안전부" in txt and "SQL 삽입" in txt and "양호" in txt
    # 항목 번호는 판마다 달라 싣지 않는다(대조 부담 제거)
    assert "SC-01" not in txt


def test_not_covered_is_not_reported_as_pass():
    """규칙이 없는 항목을 '양호'로 찍으면 안 한 점검을 했다고 쓰는 셈이다."""
    std = standards.get("efs")
    rows = standards.coverage(std, {}, available=frozenset({"CWE-89"}))
    by = {r["code"]: r["verdict"] for r in rows}
    assert by["sql-injection"] == standards.PASS            # 규칙 있음 · 탐지 0
    assert by["directory-indexing"] == standards.NOT_COVERED  # 정적 분석 대상 아님
    assert by["csrf"] == standards.NOT_COVERED               # CWE 는 있으나 규칙 없음

    rows = standards.coverage(std, {"CWE-89": 2}, available=frozenset({"CWE-89"}))
    assert next(r for r in rows if r["code"] == "sql-injection")["verdict"] == standards.VIOLATED


def test_efs_standard_is_registered():
    std = standards.get("efs")
    assert std is not None and not std.show_code
    names = {i.name for i in std.items}
    assert {"SQL 인젝션", "크로스사이트 스크립팅", "쿠키 변조"} <= names


def _seed_multi(c: Client) -> int:
    from cpguard.web.models import Scan
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("a/i.js", "app.get('/x',(req,res)=>{eval(req.query.q)})")
    c.post("/scan/", {"archive": SimpleUploadedFile("multi.zip", buf.getvalue()),
                      "standards": ["mois", "efs", "없는기준"]}, SERVER_NAME="127.0.0.1")
    return Scan.objects.order_by("-id").first().pk


def test_standards_chosen_at_scan_time_drive_the_deliverables():
    """진단 전에 고른 기준이 저장되고, ?std= 없이도 산출물에 그대로 적용돼야 한다."""
    from cpguard.web.models import Scan
    c = Client()
    pk = _seed_multi(c)
    assert Scan.objects.get(pk=pk).standard_ids == ["mois", "efs"]   # 모르는 값은 버린다

    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(
        c.get(f"/scan/{pk}/export.xlsx", SERVER_NAME="127.0.0.1").content))
    ws = wb["점검항목 결과"]
    assert [c.value for c in ws[1]][0] == "분류"          # 여러 기준 → 분류 열
    groups = {r[0] for r in ws.iter_rows(min_row=2, values_only=True) if r[0]}
    assert len(groups) == 2

    from pypdf import PdfReader
    txt = "".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(
        c.get(f"/scan/{pk}/report.pdf", SERVER_NAME="127.0.0.1").content)).pages)
    assert "전자금융감독규정" in txt and "행정안전부" in txt and "진단 대상 아님" in txt


def test_en_deliverables_contain_no_korean():
    """영문 산출물에 한글이 새면 해외 제출본이 못 쓰게 된다.

    기준 근거 문서·표 머리글처럼 서버에서 만드는 문자열은 클라이언트 사전이 손대지
    못한다 — 실제로 std.source 가 한국어 그대로 실리고 있었다."""
    import re

    c = Client()
    pk = _seed_multi(c)
    han = re.compile(r"[가-힣]")

    from pypdf import PdfReader
    pdf = c.get(f"/scan/{pk}/report.pdf?lang=en&std=mois,efs,owasp,cwe",
                SERVER_NAME="127.0.0.1").content
    txt = "".join((p.extract_text() or "") for p in PdfReader(io.BytesIO(pdf)).pages)
    assert not han.search(txt), f"PDF(en) 한글 잔존: {sorted(set(han.findall(txt)))[:20]}"

    from docx import Document
    doc = Document(io.BytesIO(c.get(f"/scan/{pk}/report.docx?lang=en&std=mois,efs",
                                    SERVER_NAME="127.0.0.1").content))
    parts = ([p.text for p in doc.paragraphs]
             + [cl.text for t in doc.tables for r in t.rows for cl in r.cells])
    joined = " ".join(parts)
    assert not han.search(joined), f"Word(en) 한글 잔존: {sorted(set(han.findall(joined)))[:20]}"

    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(
        c.get(f"/scan/{pk}/export.xlsx?lang=en&std=mois,efs", SERVER_NAME="127.0.0.1").content))
    cells = " ".join(str(cl.value) for ws in wb.worksheets
                     for row in ws.iter_rows() for cl in row if cl.value)
    assert not han.search(cells), f"xlsx(en) 한글 잔존: {sorted(set(han.findall(cells)))[:20]}"


def test_a_draft_standard_announces_itself():
    """항목 목록이 공식 문서로 확인되지 않은 기준은 화면과 산출물에서 그렇다고 말해야 한다.

    점검표는 빠진 항목이 조용히 없는 게 가장 위험하다.
    """
    from cpguard.standards import STANDARDS

    mobile = STANDARDS["mobile"]
    assert mobile.draft_note, "초안 경고가 있어야 한다"
    assert len(mobile.items) == 24
    assert all(it.cwes for it in mobile.items), "모든 항목에 CWE 매핑이 있어야 한다"

    for other in ("mois", "efs", "owasp", "cwe"):
        assert not STANDARDS[other].draft_note, f"{other} 는 초안이 아니다"
