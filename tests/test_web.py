"""대시보드 통합 테스트 — 업로드 → 안전해제 → 스캔 → 결과 렌더."""
import io
import os
import tempfile
import zipfile
from pathlib import Path

import pytest

os.environ["CPGUARD_HOME"] = tempfile.mkdtemp(prefix="cpguard_test_")
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "cpguard.web.settings")

import django  # noqa: E402

django.setup()
from django.conf import settings  # noqa: E402
from django.core.management import call_command  # noqa: E402
from django.test import Client  # noqa: E402

# 테스트 클라이언트는 host 헤더로 'testserver' 를 쓴다
settings.ALLOWED_HOSTS = ["testserver", "127.0.0.1", "localhost"]


@pytest.fixture(scope="module", autouse=True)
def _db():
    call_command("migrate", verbosity=0, interactive=False)


def _zip_bytes(files: dict[str, str]) -> io.BytesIO:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in files.items():
            z.writestr(name, data)
    buf.seek(0)
    buf.name = "project.zip"
    return buf


def test_index_loads():
    assert Client().get("/").status_code == 200


def test_upload_scan_and_detail():
    c = Client()
    z = _zip_bytes({
        "app/server.js": "app.get('/p', function(req,res){ const h=req.query.h; child_process.exec(h); });",
        "app/safe.js": "function s(req){ const h=shellQuote(req.query.h); child_process.exec(h); }",
    })
    r = c.post("/scan/", {"archive": z}, follow=True)
    assert r.status_code == 200
    body = r.content.decode("utf-8")
    assert "js.command-injection" in body      # 취약 파일은 탐지
    assert "CWE-78" in body
    assert body.count("js.command-injection") >= 1
    assert "req.query.h" in body                # 흐름 경로가 화면에 보인다


def test_zip_slip_upload_rejected():
    c = Client()
    z = _zip_bytes({"../evil.js": "pwned"})
    r = c.post("/scan/", {"archive": z}, follow=True)
    assert "안전하지 않은 아카이브" in r.content.decode("utf-8")


def test_non_zip_rejected():
    c = Client()
    f = io.BytesIO(b"not a zip")
    f.name = "x.txt"
    r = c.post("/scan/", {"archive": f}, follow=True)
    assert "zip 파일만" in r.content.decode("utf-8")


def test_cross_file_interprocedural_detection():
    """서로 다른 파일에 걸친 흐름도 탐지해야 한다."""
    c = Client()
    z = _zip_bytes({
        "lib/input.js": "function readInput(req){ return req.query.cmd; }\nmodule.exports = readInput;",
        "lib/run.js": "function runIt(v){ child_process.exec(v); }\nmodule.exports = runIt;",
        "app.js": "function handler(req,res){ const c = readInput(req); runIt(c); }",
    })
    r = c.post("/scan/", {"archive": z}, follow=True)
    body = r.content.decode("utf-8")
    assert "js.command-injection" in body


# ---------- 감사 작업대 ----------

def _seed_scan(c: Client) -> int:
    z = _zip_bytes({
        "app/server.js": "app.get('/p', function(req,res){ const h=req.query.h; child_process.exec(h); });",
        "app/conf.py": 'password = "hunter2secret"\n',
    })
    c.post("/scan/", {"archive": z}, follow=True)
    from cpguard.web.models import Scan
    return Scan.objects.first().pk


def test_workbench_renders_panes_and_data():
    c = Client()
    pk = _seed_scan(c)
    body = c.get(f"/scan/{pk}/").content.decode("utf-8")
    assert 'id="tree"' in body and 'id="code"' in body and 'id="detail"' in body
    assert 'id="data-findings"' in body and 'id="data-sources"' in body
    # 다운로드 링크는 download 속성이 있어야 한다: 네이티브 WebView2 가 PDF 를 창 안에서
    # 열어(뷰어 탈취) 먹통 되는 것을 막고 실제 다운로드로 보낸다.
    for url in (f"/scan/{pk}/export.xlsx", f"/scan/{pk}/guide.pdf"):
        i = body.index(f'href="{url}"')
        assert "download" in body[max(0, i - 60):i]
    # 합본 보고서는 여러 프로젝트를 묶는 산출물이라 검토 화면이 아니라 리포트에서 만든다
    assert f"/scan/{pk}/report.pdf" not in body
    assert "/reports/" in body


def test_sources_are_stored_for_code_viewer():
    c = Client()
    pk = _seed_scan(c)
    from cpguard.web.models import Scan
    scan = Scan.objects.get(pk=pk)
    assert scan.sources, "코드 뷰어가 쓸 원본이 보관돼야 한다"
    assert any(p.endswith("server.js") for p in scan.sources)


def test_script_tag_in_source_cannot_break_out():
    """소스에 </script> 가 있어도 페이지 스크립트를 깨뜨리면 안 된다."""
    c = Client()
    z = _zip_bytes({
        "a.js": "// </script><img src=x onerror=alert(1)>\n"
                "app.get('/p', function(req,res){ child_process.exec(req.query.h); });",
    })
    c.post("/scan/", {"archive": z}, follow=True)
    from cpguard.web.models import Scan
    body = c.get(f"/scan/{Scan.objects.first().pk}/").content.decode("utf-8")
    assert "</script><img" not in body          # 원문 그대로 새어나오면 안 됨
    assert "\u003C" in body or "\u003c" in body  # json_script 가 이스케이프


def test_audit_state_persists():
    c = Client()
    pk = _seed_scan(c)
    r = c.post(f"/scan/{pk}/audit/", {"index": "0", "status": "confirmed"})
    assert r.json()["ok"] is True
    from cpguard.web.models import Scan
    assert Scan.objects.get(pk=pk).audit["0"] == "confirmed"


def test_audit_note_saves_as_plaintext():
    c = Client()
    pk = _seed_scan(c)
    from cpguard.web.models import Scan
    payload = "오탐 근거: 테스트 코드 <script>alert(1)</script>"
    r = c.post(f"/scan/{pk}/note/", {"index": "0", "note": payload})
    assert r.json()["ok"] is True
    # 평문 그대로 저장(렌더는 프런트에서 escape) — 저장 단계에서 변형하지 않는다
    assert Scan.objects.get(pk=pk).audit_notes["0"] == payload
    # 상세/작업대 응답이 note 를 실어 준다
    d = c.get(f"/scan/{pk}/api/finding/0").json()
    assert d["finding"]["audit_note"] == payload
    # 빈 값으로 저장하면 삭제
    c.post(f"/scan/{pk}/note/", {"index": "0", "note": "  "})
    assert "0" not in Scan.objects.get(pk=pk).audit_notes


def test_audit_rejects_unknown_status():
    c = Client()
    pk = _seed_scan(c)
    r = c.post(f"/scan/{pk}/audit/", {"index": "0", "status": "hacked"})
    assert r.status_code == 400


def test_audit_requires_post():
    c = Client()
    pk = _seed_scan(c)
    assert c.get(f"/scan/{pk}/audit/").status_code == 405


def test_csv_export():
    c = Client()
    pk = _seed_scan(c)
    r = c.get(f"/scan/{pk}/export.csv")
    assert r.status_code == 200
    text = r.content.decode("utf-8-sig")
    assert "위험도" in text and "CWE" in text
    assert "js.command-injection" in text or "secret.hardcoded-password" in text


def test_downloads_are_no_store():
    """다운로드 응답은 캐시 금지여야 한다.

    SQLite 는 삭제된 pk(rowid)를 재사용한다. 네이티브 WebView2 가 /scan/<pk>/export.*
    GET 을 캐시하면, 같은 pk 를 다른 프로젝트 스캔이 재획득했을 때 이전 프로젝트의
    파일이 내려간다(현장 재현 버그). never_cache 로 이를 막는다.
    """
    c = Client()
    pk = _seed_scan(c)
    for url in (f"/scan/{pk}/export.xlsx", f"/scan/{pk}/export.csv",
                f"/scan/{pk}/sarif/", f"/scan/{pk}/report.pdf", f"/scan/{pk}/guide.pdf"):
        r = c.get(url)
        assert r.status_code == 200, url
        assert "no-store" in r.headers.get("Cache-Control", ""), url


def test_scan_progress_status_and_page():
    """진행 상태 API 와 진행 화면 렌더 (백그라운드 잡 머신)."""
    import time as _t

    from cpguard.web import views
    c = Client()
    jid = "testjob_" + os.urandom(4).hex()
    views._job_set(jid, status="running", phase="parse", done=3, total=10,
                   findings=2, name="proj.zip", started=_t.time())

    d = c.get(f"/scan/progress/{jid}/status").json()
    assert d["status"] == "running" and d["total"] == 10 and d["phase"] == "parse"
    assert d["done"] == 3 and d["findings"] == 2

    page = c.get(f"/scan/progress/{jid}/")
    assert page.status_code == 200
    body = page.content.decode("utf-8")
    assert 'id="steps"' in body and 'id="log"' in body   # 단계 체크리스트·로그 존재
    assert 'id="goto"' in body                            # 프로젝트로 이동 버튼

    assert c.get("/scan/progress/does-not-exist/status").status_code == 404


def test_run_scan_job_creates_scan():
    """백그라운드 잡 함수가 압축해제→스캔→Scan 생성까지 하고 done/pk 를 남긴다."""
    from pathlib import Path

    from cpguard.web import views
    z = _zip_bytes({"a.js": "app.get('/p',function(req,res){child_process.exec(req.query.h);});"})
    workdir = Path(tempfile.mkdtemp(prefix="cpguard_job_"))
    (workdir / "upload.zip").write_bytes(z.getvalue())
    jid = "job_" + os.urandom(4).hex()
    views._job_set(jid, status="running", name="a.zip", started=0)
    views._run_scan_job(jid, workdir, "a.zip", False, "")
    job = views._job_get(jid)
    assert job["status"] == "done" and isinstance(job.get("pk"), int)
    assert job["findings"] >= 1


def test_settings_save_mask_and_clear():
    """설정 저장 → 마스킹 표시 → 삭제. 환경 오염은 끝에 정리."""
    import os as _os

    from cpguard.web import config as appcfg
    c = Client()
    try:
        r = c.post("/settings/", {"ANTHROPIC_API_KEY": "sk-ant-DUMMY1234567890abcd"})
        assert r.status_code in (302, 200)
        assert appcfg.load().get("ANTHROPIC_API_KEY") == "sk-ant-DUMMY1234567890abcd"
        body = c.get("/settings/").content.decode("utf-8")
        assert "설정됨" in body and "•" in body          # 마스킹 표시
        assert "sk-ant-DUMMY1234567890abcd" not in body  # 원본 노출 안 함
        # 삭제
        c.post("/settings/", {"clear_ANTHROPIC_API_KEY": "1"})
        assert "ANTHROPIC_API_KEY" not in appcfg.load()
    finally:
        appcfg.save({})
        _os.environ.pop("ANTHROPIC_API_KEY", None)


def test_guide_page_renders_markdown():
    c = Client()
    r = c.get("/guide/")
    assert r.status_code == 200
    body = r.content.decode("utf-8")
    assert "사용 가이드" in body
    assert "<h2>" in body and "<ol>" in body and "<pre>" in body   # 마크다운→HTML
    # 프로바이더 키 발급 가이드가 설정에 노출되는지
    s = c.get("/settings/").content.decode("utf-8")
    assert "발급 가이드" in s and 'id="guide-gemini"' in s


def test_settings_saves_workspace_id():
    """워크스페이스 ID(신원 연동 Claude 키용) 저장·적용·삭제."""
    import os as _os

    from cpguard.web import config as appcfg
    c = Client()
    try:
        c.post("/settings/", {"ANTHROPIC_WORKSPACE_ID": "wrkspc_test123"})
        assert appcfg.load().get("ANTHROPIC_WORKSPACE_ID") == "wrkspc_test123"
        assert _os.environ.get("ANTHROPIC_WORKSPACE_ID") == "wrkspc_test123"
        c.post("/settings/", {"clear_ANTHROPIC_WORKSPACE_ID": "1"})
        assert "ANTHROPIC_WORKSPACE_ID" not in appcfg.load()
    finally:
        appcfg.save({})
        _os.environ.pop("ANTHROPIC_WORKSPACE_ID", None)


def test_secrets_only_skips_dataflow():
    """secrets_only 스캔은 데이터 흐름 축을 건너뛰고 패턴만 탐지한다."""
    from cpguard.scanner import scan_path
    d = tempfile.mkdtemp(prefix="cpguard_so_")
    with open(os.path.join(d, "a.py"), "w", encoding="utf-8") as f:
        f.write('API_KEY = "AKIAIOSFODNN7EXAMPLE"\n'
                'q = "SELECT * FROM t WHERE x=" + user\n')
    findings, report = scan_path(d, secrets_only=True)
    assert report.scanned == 0                       # 소스 파싱 축 생략
    cats = {getattr(f, "category", "flow") for f in findings}
    assert "flow" not in cats                         # 데이터 흐름 finding 없음
    assert any(f.category in ("secret", "pii", "config", "infra", "hygiene") for f in findings)


def test_dashboard_shows_stats():
    c = Client()
    body = c.get("/").content.decode("utf-8")
    assert "위험도 분포" in body and "상위 탐지 규칙" in body   # 대시보드 위젯
    assert "총 탐지" in body                                    # 상태 타일


def test_compare_and_reports_pages():
    c = Client()
    assert c.get("/compare/").status_code == 200
    assert c.get("/reports/").status_code == 200


def test_settings_saves_model_override():
    import os as _os

    from cpguard.web import config as appcfg
    c = Client()
    try:
        c.post("/settings/", {"model_gemini": "gemini-3.6-pro"})
        assert appcfg.model_for("gemini") == "gemini-3.6-pro"
        # 비우면 기본으로 복귀
        c.post("/settings/", {"model_gemini": ""})
        assert appcfg.model_for("gemini") is None
    finally:
        appcfg.save({})
        _os.environ.pop("GEMINI_API_KEY", None)


def test_pdf_report_and_guide_download():
    """합본 보고서·조치가이드 PDF 다운로드 (유효한 PDF)."""
    c = Client()
    pk = _seed_scan(c)
    for url in (f"/scan/{pk}/report.pdf", f"/scan/{pk}/guide.pdf"):
        r = c.get(url)
        assert r.status_code == 200
        assert r["Content-Type"] == "application/pdf"
        assert r.content[:5] == b"%PDF-"       # PDF 매직
        assert len(r.content) > 1500


def test_upload_accepts_model_override():
    from cpguard.web import views
    from pathlib import Path
    z = _zip_bytes({"a.js": "app.get('/p',function(req,res){child_process.exec(req.query.h);});"})
    workdir = Path(tempfile.mkdtemp(prefix="cpguard_m_"))
    (workdir / "upload.zip").write_bytes(z.getvalue())
    jid = "jm_" + os.urandom(4).hex()
    views._job_set(jid, status="running", name="a.zip", started=0)
    # model 인자를 받아도 정상 완료 (트리아지 off 라 실제 호출은 없음)
    views._run_scan_job(jid, workdir, "a.zip", False, "", False, "claude-haiku-4-5")
    assert views._job_get(jid)["status"] == "done"


def test_finding_rows_and_scale_apis():
    """FindingRow 적재 + 집계/페이지네이션 API (대량 탐지 서버측 질의)."""
    c = Client()
    pk = _seed_scan(c)
    from cpguard.web.models import FindingRow
    assert FindingRow.objects.filter(scan_id=pk).count() > 0

    s = c.get(f"/scan/{pk}/api/summary").json()
    assert s["total"] > 0 and s["severity"] and s["top_rules"]
    # 프로젝트 홈 차트가 쓰는 확장 집계 필드
    assert "by_category" in s and "top_cwe" in s and "by_verdict" in s

    d = c.get(f"/scan/{pk}/api/findings?size=1&page=1").json()
    assert d["size"] == 1 and len(d["rows"]) <= 1 and d["total"] >= 1
    assert {"severity", "rule_id", "file", "line"} <= set(d["rows"][0].keys())

    hi = c.get(f"/scan/{pk}/api/findings?severity=high").json()
    assert all(r["severity"] == "high" for r in hi["rows"])


# ---- 다건 업로드 · 배치 · 포트폴리오 (담당자 300+ 프로젝트 시나리오) ----

def _named_zip(name: str, files: dict[str, str]) -> io.BytesIO:
    z = _zip_bytes(files)
    z.name = name
    return z


def test_split_batch_zip_detects_nested():
    from cpguard.web import views
    inner = _zip_bytes({"a/x.js": "eval(req.query.q)"}).getvalue()
    outer = io.BytesIO()
    with zipfile.ZipFile(outer, "w") as z:
        z.writestr("proj-one.zip", inner)
        z.writestr("proj-two.zip", inner)
    p = Path(tempfile.mkdtemp()) / "batch.zip"
    p.write_bytes(outer.getvalue())
    parts = views._split_batch_zip(p)
    assert parts is not None and len(parts) == 2
    assert {n for n, _ in parts} == {"proj-one.zip", "proj-two.zip"}
    # 일반 프로젝트 zip(내부 zip 없음)은 None → 단일 프로젝트로 취급
    plain = Path(tempfile.mkdtemp()) / "plain.zip"
    plain.write_bytes(_zip_bytes({"a/x.js": "x=1"}).getvalue())
    assert views._split_batch_zip(plain) is None


def test_multi_file_upload_creates_batch():
    from cpguard.web.models import Scan
    c = Client()
    before = Scan.objects.count()
    z1 = _named_zip("alpha.zip", {"a/i.js": "app.get('/x',(req,res)=>{eval(req.query.q)})"})
    z2 = _named_zip("beta.zip", {"b/m.js": "app.get('/y',(req,res)=>{eval(req.query.p)})"})
    r = c.post("/scan/", {"archive": [z1, z2]})
    assert r.status_code == 302 and "/scan/batch/" in r["Location"]
    assert Scan.objects.count() - before == 2
    bid = r["Location"].split("/batch/")[1].split("/")[0]
    st = c.get(f"/scan/batch/{bid}/status").json()
    assert st["total"] == 2 and st["complete"] is True and st["done"] == 2


def test_nested_batch_zip_expands_to_projects():
    from cpguard.web.models import Scan
    c = Client()
    inner = _zip_bytes({"a/x.js": "app.get('/x',(req,res)=>{eval(req.query.q)})"}).getvalue()
    outer = io.BytesIO()
    with zipfile.ZipFile(outer, "w") as z:
        z.writestr("gamma.zip", inner)
        z.writestr("delta.zip", inner)
    outer.name = "container.zip"
    outer.seek(0)
    before = Scan.objects.count()
    r = c.post("/scan/", {"archive": outer})
    assert r.status_code == 302 and "/scan/batch/" in r["Location"]
    assert Scan.objects.count() - before == 2


def test_portfolio_page_and_severity_columns():
    from cpguard.web.models import Scan
    c = Client()
    # 스캔 하나 만들어 위험도 컬럼이 채워지는지
    c.post("/scan/", {"archive": _named_zip("portf.zip",
           {"a/i.js": "app.get('/x',(req,res)=>{eval(req.query.q)})"})})
    s = Scan.objects.order_by("-id").first()
    assert (s.sev_critical + s.sev_high + s.sev_medium + s.sev_low + s.sev_info) == s.finding_count
    r = c.get("/projects/")
    assert r.status_code == 200
    body = r.content.decode("utf-8")
    assert "포트폴리오" in body


def test_portfolio_export_zip_of_deliverables():
    from cpguard.web.models import Scan
    c = Client()
    c.post("/scan/", {"archive": _named_zip("exp.zip",
           {"a/i.js": "app.get('/x',(req,res)=>{eval(req.query.q)})"})})
    s = Scan.objects.order_by("-id").first()
    r = c.get(f"/projects/export.zip?ids={s.pk}&kind=both")
    assert r.status_code == 200 and r["Content-Type"] == "application/zip"
    zf = zipfile.ZipFile(io.BytesIO(r.content))
    names = zf.namelist()
    assert any(n.endswith("_report.pdf") for n in names)
    assert any(n.endswith("_analysis-sheet.xlsx") for n in names)


# ---- 다운로드 파일명 (긴 프로젝트명·한글) ----

def test_download_name_truncates_and_sanitizes():
    from cpguard.web.views import _download_name
    n = _download_name("a" * 300, "소스코드_취약점진단_분석목록표", "xlsx")
    assert len(n.encode("utf-8")) <= 255          # OS 파일명 한계
    assert n.endswith("_소스코드_취약점진단_분석목록표.xlsx")
    # 경로 구분자·따옴표·개행은 헤더/파일명을 깨므로 제거
    bad = _download_name('x"y\r\nz/w', "진단결과보고서", "pdf")
    assert not set(bad) & set('"/\\r\n')
    assert _download_name("", "진단결과보고서", "pdf").startswith("scan_")


def test_download_headers_are_ascii_rfc5987():
    """한글 파일명을 헤더에 그대로 넣으면 Django 가 RFC 2047 로 통째 인코딩해
    브라우저에서 파일명이 깨진다. filename* 로 실어 헤더는 ASCII 로 유지한다."""
    from urllib.parse import unquote
    c = Client()
    z = _zip_bytes({"a/i.js": "app.get('/x',(req,res)=>{eval(req.query.q)})"})
    z.name = ("p" * 200) + ".zip"
    c.post("/scan/", {"archive": z})
    from cpguard.web.models import Scan
    pk = Scan.objects.order_by("-id").first().pk
    for url in (f"/scan/{pk}/report.pdf", f"/scan/{pk}/export.xlsx", f"/scan/{pk}/export.csv"):
        cd = c.get(url)["Content-Disposition"]
        assert cd.isascii() and not cd.startswith("=?"), f"{url}: {cd[:40]}"
        assert "filename*=UTF-8''" in cd
        assert len(unquote(cd.split("filename*=UTF-8''")[1]).encode("utf-8")) <= 255


def test_installer_spec_bundles_every_grammar():
    """설치본 스펙이 모든 tree-sitter 문법을 담는지.

    손으로 적은 목록을 쓰던 시절 언어 7개가 빠져 있었다 — 개발 환경에서는 잘 되고
    설치본에서만 그 언어가 조용히 동작하지 않아 알아채기 어렵다."""
    import re
    root = Path(__file__).resolve().parent.parent
    spec = (root / "packaging" / "cpguard.spec").read_text(encoding="utf-8")
    loader = (root / "cpguard" / "parse" / "loader.py").read_text(encoding="utf-8")
    grammars = set(re.findall(r"^import (tree_sitter\w*)", loader, re.M))
    assert len(grammars) >= 12
    # 스펙은 목록을 하드코딩하지 않고 loader 에서 뽑아야 한다
    assert "loader.py" in spec and r"tree_sitter\w*" in spec
    assert not re.search(r'"tree_sitter_javascript",\s*"tree_sitter_typescript"', spec)


def test_word_report_downloads():
    """Word 산출물 — 담당자가 고쳐 쓰는 원본."""
    import io

    from docx import Document
    c = Client()
    pk = _seed_scan(c)
    r = c.get(f"/scan/{pk}/report.docx?std=mois", SERVER_NAME="127.0.0.1")
    assert r.status_code == 200
    assert "wordprocessingml" in r["Content-Type"]
    d = Document(io.BytesIO(r.content))
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    assert "1. 진단 개요" in heads and "3. 진단 항목" in heads and "5. 종합 의견" in heads
    # 고른 기준의 점검표가 실려야 한다
    cells = {c.text for t in d.tables for row in t.rows for c in row.cells}
    assert "SQL 삽입" in cells and "양호" in cells


def test_pdf_report_is_grayscale():
    """산출물은 검은 글자 + 옅은 회색으로 통일한다.

    진단 보고서는 흑백 출력·복사본으로 돌아다니고 발주처 문서 양식에 얹히는 일이 많다.
    색으로만 구분되는 정보가 있으면 그 과정에서 사라진다."""
    import re

    from pypdf import PdfReader
    c = Client()
    pk = _seed_scan(c)
    data = c.get(f"/scan/{pk}/report.pdf?std=mois", SERVER_NAME="127.0.0.1").content
    colors = set()
    for page in PdfReader(io.BytesIO(data)).pages:
        ops = page.get_contents().get_data().decode("latin-1", "replace")
        for m in re.finditer(r"([\d.]+) ([\d.]+) ([\d.]+) (?:rg|RG)", ops):
            colors.add(tuple(round(float(m.group(i)), 3) for i in (1, 2, 3)))
    assert colors, "색 연산자를 하나도 못 찾았다 — 검사가 헛돌고 있다"
    non_gray = [c for c in colors
                if abs(c[0] - c[1]) > 0.02 or abs(c[1] - c[2]) > 0.02]
    assert not non_gray, f"회색이 아닌 색이 남아 있다: {sorted(non_gray)}"


def test_readme_counts_match_reality():
    """README 가 광고하는 숫자와 실제가 어긋나면 첫인상부터 신뢰를 잃는다.

    규칙을 추가할 때마다 손으로 고치다 보니 실제로 두 번 어긋났다."""
    import re

    from cpguard.parse import loader
    from cpguard.taint.spec import load_rules
    rules, exts = len(load_rules()), len(loader.SUPPORTED_EXTENSIONS)
    root = Path(__file__).resolve().parent.parent
    for name in ("README.md", "README.ko.md"):
        txt = (root / name).read_text(encoding="utf-8")
        badge = re.search(r"badge/taint%20rules-(\d+)-", txt)
        assert badge and int(badge.group(1)) == rules, f"{name} 배지: {badge and badge.group(1)} != {rules}"
        assert f"**{rules} rules**" in txt or f"**규칙 {rules}개**" in txt, f"{name} 본문 규칙 수"
        assert f"{exts} file extensions" in txt or f"확장자 {exts}종" in txt, f"{name} 확장자 수"


def test_audit_moves_the_open_count_not_the_detection_total():
    """감사 상태가 조치대상 건수를 줄인다 — 탐지 총계는 스캔이 찾은 사실이라 그대로다.

    화면 숫자와 합본 보고서의 '최종 조치대상'이 같은 기준을 써야 제출물과 어긋나지 않는다.
    """
    from cpguard.report.consolidated import AUDIT_REASON
    from cpguard.web.models import CLOSED_AUDIT, Scan

    # 조치대상에서 빠지는 상태는 산출물 쪽 정의와 같아야 한다
    assert set(CLOSED_AUDIT) == set(AUDIT_REASON)

    c = Client()
    pk = _seed_scan(c)
    scan = Scan.objects.get(pk=pk)
    total = scan.finding_count
    assert scan.open_count == total          # 감사 전에는 전부 조치대상

    scan.set_audit(0, "false_positive")
    scan.set_audit(1, "confirmed")           # 취약 확정은 조치대상으로 남는다
    scan = Scan.objects.get(pk=pk)

    assert scan.finding_count == total
    assert scan.open_count == total - 1
    assert scan.audit_summary["false_positive"] == 1
    assert scan.audit_summary["confirmed"] == 1
    assert scan.audit_summary["unaudited"] == total - 2
    assert sum(scan.open_severity_counts.values()) == total - 1


def test_audit_rejects_an_index_that_is_not_a_finding():
    """범위 밖 index 를 받아주면 조치대상 집계가 실재하지 않는 항목만큼 어긋난다."""
    from cpguard.web.models import Scan

    c = Client()
    pk = _seed_scan(c)
    before = Scan.objects.get(pk=pk).open_count
    r = c.post(f"/scan/{pk}/audit/", {"index": 999999, "status": "fixed"})
    assert r.status_code == 400
    assert Scan.objects.get(pk=pk).open_count == before   # 집계가 흔들리지 않는다


def test_workbench_ships_the_live_count_hooks():
    """감사 판정이 새로고침 없이 헤더 숫자에 반영되려면 훅과 초기값이 함께 나가야 한다."""
    c = Client()
    pk = _seed_scan(c)
    html = c.get(f"/scan/{pk}/", SERVER_NAME="127.0.0.1").content.decode()
    assert 'id="c-open"' in html and 'id="c-sev"' in html
    assert 'data-open-counts' in html          # 위험도 초기값
    assert "applyAuditDelta" in html           # 판정 시 증감


def _nested_zip(inner: dict[str, dict[str, str]], extra: dict[str, str] | None = None) -> io.BytesIO:
    """zip 안에 zip 을 넣은 아카이브. extra 는 zip 이 아닌 동봉 파일."""
    outer = io.BytesIO()
    with zipfile.ZipFile(outer, "w") as zf:
        for name, files in inner.items():
            zf.writestr(name, _zip_bytes(files).getvalue())
        for name, text in (extra or {}).items():
            zf.writestr(name, text)
    outer.seek(0)
    return outer


def test_a_project_carrying_zip_dependencies_is_one_project(tmp_path):
    """yarn PnP 처럼 의존성을 zip 으로 들고 다니는 프로젝트를 쪼개면 안 된다.

    실제로 5GB 프로젝트 하나가 1,762건의 '프로젝트'로 갈라졌다.
    """
    from cpguard.web.views import _split_batch_zip

    z = tmp_path / "proj.zip"
    z.write_bytes(_nested_zip(
        {".yarn/cache/@babel-parser-npm-7.29.2.zip": {"index.js": "module.exports = 1;\n"},
         ".yarn/cache/@babel-runtime-npm-7.29.2.zip": {"index.js": "module.exports = 2;\n"}},
        extra={"src/app.js": "app.get('/x', (req,res) => eval(req.query.c));\n"},
    ).getvalue())
    assert _split_batch_zip(z) is None


def test_an_archive_of_project_zips_is_still_a_batch(tmp_path):
    from cpguard.web.views import _split_batch_zip

    z = tmp_path / "batch.zip"
    z.write_bytes(_nested_zip({
        "alpha.zip": {"a.js": "const x = 1;\n"},
        "beta.zip": {"b.js": "const y = 2;\n"},
    }, extra={"__MACOSX/._alpha.zip": "junk"}).getvalue())     # 부산물은 무시한다
    parts = _split_batch_zip(z)
    assert parts is not None and sorted(n for n, _ in parts) == ["alpha.zip", "beta.zip"]


def test_a_single_inner_zip_is_not_a_batch(tmp_path):
    from cpguard.web.views import _split_batch_zip

    z = tmp_path / "one.zip"
    z.write_bytes(_nested_zip({"only.zip": {"a.js": "const x = 1;\n"}}).getvalue())
    assert _split_batch_zip(z) is None


def test_bulk_delete_removes_only_the_checked_scans():
    """배치를 잘못 돌리면 수백 건이 쌓인다 — 골라서 지울 수 있어야 한다."""
    from cpguard.web.models import Scan

    c = Client()
    keep, drop1, drop2 = _seed_scan(c), _seed_scan(c), _seed_scan(c)
    before = Scan.objects.count()          # 다른 테스트가 남긴 스캔이 있을 수 있다

    r = c.post("/scan/delete-many/", {"pk": [drop1, drop2]}, follow=True)
    assert r.status_code == 200
    assert Scan.objects.count() == before - 2
    assert Scan.objects.filter(pk=keep).exists()
    assert not Scan.objects.filter(pk__in=[drop1, drop2]).exists()


def test_cancelling_a_scan_marks_it_cancelled():
    """중단은 진행 콜백에서 걸린다 — 요청은 플래그만 세우고 즉시 응답한다."""
    from cpguard.web import views

    job_id = "testjob"
    views._job_set(job_id, status="running", name="x.zip")
    try:
        r = Client().post(f"/scan/progress/{job_id}/cancel")
        assert r.status_code == 200 and r.json()["ok"]
        assert views._cancelled(job_id)
    finally:
        views._JOBS.pop(job_id, None)


def test_a_batch_cancel_reaches_its_queued_jobs():
    from cpguard.web import views

    views._job_set("j1", status="queued", batch_id="b1")
    views._batch_set("b1", job_ids=["j1"], total=1)
    try:
        assert Client().post("/scan/batch/b1/cancel").status_code == 200
        assert views._cancelled("j1")          # 대기 중이던 항목도 시작하지 않는다
    finally:
        views._JOBS.pop("j1", None)
        views._BATCHES.pop("b1", None)


def test_running_scans_feed_the_return_banner():
    """다른 화면에 있어도 진행 중인 진단으로 돌아갈 수 있어야 한다."""
    from cpguard.web import views

    views._job_set("j2", status="running", name="live.zip")
    try:
        running = Client().get("/scan/running").json()["running"]
        assert any(r["name"] == "live.zip" and "/scan/progress/j2/" in r["url"] for r in running)
    finally:
        views._JOBS.pop("j2", None)


def test_a_rescan_inherits_the_previous_verdicts():
    """재점검에서 지난번 오탐을 다시 판정하지 않는다 — 지문이 같으면 판정을 잇는다."""
    from cpguard.web.models import Scan

    c = Client()
    first = _seed_scan(c)
    s1 = Scan.objects.get(pk=first)
    s1.set_audit(0, "false_positive")
    s1.set_audit_note(0, "테스트 픽스처라 제외")

    second = _seed_scan(c)                     # 같은 내용 재업로드 = 같은 프로젝트
    s2 = Scan.objects.get(pk=second)
    assert s2.pk != s1.pk

    fp0 = s1.findings[0]["fp"]
    carried = [f for f in s2.findings if f["fp"] == fp0]
    assert carried, "지문이 같은 이슈가 있어야 한다"
    idx = str(carried[0]["id"])
    assert s2.audit.get(idx) == "false_positive"
    assert s2.audit_notes.get(idx) == "테스트 픽스처라 제외"
    assert s2.open_count == s2.finding_count - 1      # 승계된 오탐만큼 조치대상이 준다


def test_a_project_with_no_history_inherits_nothing():
    """이력이 없는 프로젝트는 승계할 것도 없다."""
    from cpguard.web.views import _carry_over_audit

    audit, notes = _carry_over_audit("존재하지-않는-프로젝트", [], Path("."))
    assert audit == {} and notes == {}
